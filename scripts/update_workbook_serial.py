#!/usr/bin/env python3
"""Update the ADM FireNode workbook to reflect the ESP32 serial reader integration."""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except Exception as exc:
    print(f"ERROR: openpyxl is required: {exc}", file=sys.stderr)
    raise SystemExit(1)

WORKBOOK = Path(__file__).resolve().parents[1] / "docs" / "workbook" / "ADM_FireNode_Implementation_Workbook.xlsx"
DATE_STR = datetime.now().strftime("%Y-%m-%d")


def update_dashboard(wb):
    ws = wb["Dashboard"]
    ws["B10"] = (
        "ESP32 USB serial reader integrated for MAIN RPi dashboard; "
        "MAIN and NODE_01 sensor data via /dev/ttyUSB0 115200; "
        "HTTP ESP32 fallback preserved; no camera/thermal/firmware changes."
    )
    ws["B11"] = (
        "Deploy updated raspi/firenode-system to .51 MAIN and validate live mode sensor cards. "
        "Confirm NODE_01 serial packets populate remote slot 1."
    )


def update_progress_tracker(wb):
    ws = wb["Progress Tracker"]
    max_row = ws.max_row
    new_rows = [
        [
            "P043", "Integration", "Integrate MAIN ESP32 USB serial data into RPi dashboard",
            "Done", "Codex", DATE_STR,
            "modules/esp32_serial_reader.py created; app.py uses serial cache; py_compile passed; config.json updated",
            "Deploy to .51 and validate with live serial traffic"
        ],
        [
            "P044", "Integration", "ESP32 serial parsing and caching per node ID",
            "Done", "Codex", DATE_STR,
            "NODE=MAIN and NODE=NODE_01 packets parsed; decorative lines ignored; thread-safe cache",
            "Validate on hardware with minicom traffic"
        ],
        [
            "P045", "Integration", "Dashboard mapping for serial sensor data (MAIN and remote nodes)",
            "Done", "Codex", DATE_STR,
            "MAIN maps to local node; NODE_01/02/03 map to remote slots 1/2/3; placeholders remain when no data",
            "Field test with NODE_01 transmitting"
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def update_issue_log(wb):
    ws = wb["Issue Log"]
    max_row = ws.max_row
    new_rows = [
        [
            "I016", "Integration", "No dashboard serial reader existed for MAIN ESP32 USB output",
            "New module modules/esp32_serial_reader.py reads /dev/ttyUSB0, parses packets, and feeds dashboard.",
            "Codex", "Resolved",
            "Serial reader starts automatically in server live mode. Configurable via esp32_serial_* keys.",
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def update_decision_log(wb):
    ws = wb["Decision Log"]
    max_row = ws.max_row
    new_rows = [
        [
            "D009", DATE_STR,
            "MAIN RPi reads ESP32 sensor data over USB serial instead of HTTP API",
            "MAIN ESP32 is physically wired by USB to .51 RPi; serial output is live and reliable; HTTP scan/fetch remains fallback",
            "Minicom confirms live serial output at 115200",
            "Serial is primary for MAIN local data and remote NODE_01 packets on MAIN dashboard; HTTP still used for remote RPi node pulls",
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def update_build_validation_matrix(wb):
    ws = wb["Build Validation Matrix"]
    max_row = ws.max_row
    new_rows = [
        [
            "ESP32 USB serial reader module",
            "Pass local compile/tests",
            "Pending hardware validation on .51",
            "In Progress",
            "modules/esp32_serial_reader.py; parsing unit test passed; py_compile passed",
            "python3 -m py_compile; manual parse_line test",
            "Thread-safe cache; configurable port/baud; decorative lines ignored.",
        ],
        [
            "Dashboard serial data overlay (MAIN + remote nodes)",
            "Pass local compile/tests",
            "Pending hardware validation on .51",
            "In Progress",
            "app.py updated: get_local_node_data and build_remote_slots overlay serial cache; HTTP fallback preserved",
            "python3 -m py_compile; deploy to .51; verify /api/server-dashboard",
            "No camera/thermal/firmware changes.",
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def update_config_tracker(wb):
    ws = wb["Config Tracker"]
    max_row = ws.max_row
    new_rows = [
        [
            "esp32_serial_enabled", "True (server live mode)", ".51 MAIN", "Serial reader integration 2026-06-02",
            "Approved", "Enables USB serial reader on MAIN RPi.",
        ],
        [
            "esp32_serial_port", "/dev/ttyUSB0", ".51 MAIN", "Hardware wiring",
            "Approved", "MAIN ESP32 USB-serial device path.",
        ],
        [
            "esp32_serial_baud", "115200", ".51 MAIN", "ESP32 firmware default",
            "Approved", "Matches ESP32 serial baud rate.",
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def main():
    if not WORKBOOK.exists():
        print(f"ERROR: Workbook not found: {WORKBOOK}", file=sys.stderr)
        raise SystemExit(2)
    wb = openpyxl.load_workbook(WORKBOOK)
    update_dashboard(wb)
    update_progress_tracker(wb)
    update_issue_log(wb)
    update_decision_log(wb)
    update_build_validation_matrix(wb)
    update_config_tracker(wb)
    wb.save(WORKBOOK)
    print(f"Workbook updated: {WORKBOOK}")


if __name__ == "__main__":
    main()
