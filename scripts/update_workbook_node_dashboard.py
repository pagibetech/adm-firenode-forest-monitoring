#!/usr/bin/env python3
"""Update the ADM FireNode workbook for the separate node dashboard GUI."""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except Exception as exc:
    print(f"ERROR: openpyxl is required: {exc}", file=sys.stderr)
    raise SystemExit(1)

WORKBOOK = Path(__file__).resolve().parents[1] / "docs" / "workbook" / "ADM_FireNode_Implementation_Workbook.xlsx"
DATE_STR = datetime.now().strftime("%Y-%m-%d")


def update_dashboard(wb):
    ws = wb["Dashboard"]
    ws["B10"] = (
        "Separate NODE GUI implemented: node_dashboard.html + node_app.js for NODE RPis only (192.168.9.52/.53/.54). "
        "NODE GUI is one-page only with local camera, sensor readings, chainsaw controls/status, alerts, recordings. "
        "MAIN server dashboard (dashboard.html + app.js) preserved unchanged. "
        "NODE simulation mode forced to live. Chainsaw controls moved to NODE dashboard. "
        "No ESP32/LoRa/thermal architecture changes."
    )


def update_progress_tracker(wb):
    ws = wb["Progress Tracker"]
    max_row = ws.max_row
    new_rows = [
        [
            "P046", "Frontend", "Create separate one-page web GUI for NODE RPis only",
            "Done", "Codex", DATE_STR,
            "node_dashboard.html and node_app.js created; app.py routes node role to node_dashboard.html",
            "Deploy to .52/.53/.54 and verify NODE GUI loads correctly"
        ],
        [
            "P047", "Frontend", "Move chainsaw controls to NODE dashboard page",
            "Done", "Codex", DATE_STR,
            "Chainsaw start/stop, sensitivity settings, audio test browser in node_dashboard.html",
            "Deploy and verify chainsaw detection from NODE GUI"
        ],
        [
            "P048", "Frontend", "Disable simulation mode for NODE RPis",
            "Done", "Codex", DATE_STR,
            "operation_mode() forced to live when current_role() == node; simulation UI removed from NODE GUI",
            "Deploy and verify NODE-only live mode on .52/.53/.54"
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def update_decision_log(wb):
    ws = wb["Decision Log"]
    max_row = ws.max_row
    new_rows = [
        [
            "D010", DATE_STR,
            "Separate NODE GUI vs shared dashboard.html",
            "NODE RPis need a simplified one-page dashboard without server-specific cards; MAIN dashboard must remain full-featured",
            "Maintainability and role-appropriate UX",
            "app.py index() routes by role; node_dashboard.html is one-page; dashboard.html unchanged for MAIN",
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def update_build_validation_matrix(wb):
    ws = wb["Build Validation Matrix"]
    max_row = ws.max_row
    new_rows = [
        [
            "Node dashboard (node_dashboard.html + node_app.js)",
            "Pass local py_compile and role routing logic",
            "Pending deployment to .52/.53/.54",
            "In Progress",
            "templates/node_dashboard.html, static/node_app.js; app.py updated",
            "python3 -m py_compile app.py modules/*.py",
            "One-page NODE GUI with camera, sensors, chainsaw, alerts, recordings.",
        ],
    ]
    for row in new_rows:
        max_row += 1
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=max_row, column=col_idx, value=value)


def main():
    if not WORKBOOK.exists():
        print(f"ERROR: Workbook not found: {WORKBOOK}", file=sys.stderr)
        raise SystemExit(2)
    wb = openpyxl.load_workbook(WORKBOOK)
    update_dashboard(wb)
    update_progress_tracker(wb)
    update_decision_log(wb)
    update_build_validation_matrix(wb)
    wb.save(WORKBOOK)
    print(f"Workbook updated: {WORKBOOK}")


if __name__ == "__main__":
    main()
