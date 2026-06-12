#!/usr/bin/env python3
"""Sync the ADM FireNode workbook with current project state (2026-06-12)."""
from pathlib import Path
import openpyxl

root = Path('/Users/macbookm1max321tb/A_Design/A_Coding/ADM Fire')
wb_path = root / 'docs/workbook/ADM_FireNode_Implementation_Workbook.xlsx'

book = openpyxl.load_workbook(wb_path)

# Sheet 1: Dashboard
ws1 = book['Dashboard']
current_stage = """ADM FireNode v4 — 2026-06-12:
- MAIN .51: CSI camera operational; dashboard LIVE mode; ESP32 USB serial integration implemented (verification pending .51 availability).
- NODE_01 .52: CSI camera operational; appears on MAIN dashboard; ESP32 bench validated; LoRa two-way confirmed.
- NODE_02 .53: CSI camera operational on local dashboard; ESP32 + LoRa hardware pending assembly.
- NODE_03 .54: CSI camera operational; UART readiness verified (/dev/serial0, serial reader connected, waiting for ESP32 wiring).
- NODE dashboard: separate one-page GUI for .52/.53/.54; chainsaw detector UX (device selection, error visibility, Starting feedback, audio visualizer); sample rate fix (device default); live device enumeration v2 (arecord -l non-zero exit now authoritative — no stale PortAudio fallback).
- Multi-node unified test firmware: 4 build environments (main/node_01/node_02/node_03) producing distinct firmware.bin files; all build SUCCESS.
- Blocker: MAIN .51 unavailable for centralized serial validation. NODE_02/NODE_03 ESP32 hardware not yet assembled.
- Next physical task: wire ESP32 to NODE_03 GPIO UART."""

next_action = """NODE_03 UART wiring (when hardware available). Deploy updated node_app to .54. Verify NODE01/NODE02 UART readiness. Update workbook to reflect all implemented features. Commits: 423fb63 (device enum v2), daf78c4 (multi-node firmware)."""

for row in ws1.iter_rows(min_row=1, max_row=20, min_col=2, max_col=2):
    for cell in row:
        if cell.value and isinstance(cell.value, str) and 'ADM FireNode' in cell.value:
            cell.value = current_stage
            print(f"Updated Dashboard Current Stage at {cell.coordinate}")
        if cell.value and isinstance(cell.value, str) and 'Separate NODE GUI' in cell.value:
            cell.value = next_action
            print(f"Updated Dashboard Next Action at {cell.coordinate}")

# Sheet 16: Progress Tracker
ws16 = book['Progress Tracker']
last_row = ws16.max_row
while last_row > 1:
    has_data = any(ws16.cell(row=last_row, column=c).value is not None for c in range(1, 7))
    if has_data:
        break
    last_row -= 1
next_row = last_row + 1

new_steps = [
    ['P033', 'Dashboard', 'Separate NODE dashboard (node_dashboard.html + node_app.js) for NODE RPis only; MAIN dashboard preserved', 'Done', 'Codex', '2026-06-09', 'node_dashboard.html; node_app.js; app.py role routing', 'NODE GUI cleanup'],
    ['P034', 'Detection', 'Chainsaw detector UX: device selection, error visibility, Starting feedback, audio visualizer', 'Done', 'Codex', '2026-06-10', 'app.py; detector.py; node_app.js', 'Field validation with real mic'],
    ['P035', 'Detection', 'Sample rate fix: detector uses device default_samplerate instead of hardcoded 16000 Hz', 'Done', 'Codex', '2026-06-10', 'detector.py _loop() updated', 'Verify on all USB mic models'],
    ['P036', 'Detection', 'Live device enumeration v2: arecord -l non-zero exit authoritative (no stale PortAudio)', 'Done', 'Codex', '2026-06-12', 'app.py; tests/test_api_devices.py', 'Field validate on RPi'],
    ['P037', 'Hardware', 'NODE03 (.54) UART readiness verified: /dev/serial0, serial reader connected', 'Done', 'Codex', '2026-06-12', 'API local_serial.connected=true', 'Wire ESP32 to GPIO UART'],
    ['P038', 'Firmware', 'Multi-node firmware builds: 4 PlatformIO environments (main/node_01/node_02/node_03)', 'Done', 'Codex', '2026-06-12', '4 firmware.bin files built; correct NODE_ID', 'Flash when ESP32 hardware arrives'],
]

for i, step in enumerate(new_steps):
    for j, val in enumerate(step):
        ws16.cell(row=next_row + i, column=j + 1, value=val)

print(f"Progress Tracker: added {len(new_steps)} steps")

# Sheet 24: Active Context Snapshot
ws24 = book['Active Context Snapshot']
updates_24 = {
    'Latest Completed Step': 'Multi-node firmware builds; live device enumeration v2 fix; NODE03 UART readiness verified. Commits: 423fb63, daf78c4.',
    'Current Active Milestone': 'NODE03 UART physical wiring (pending hardware); MAIN .51 serial validation (pending .51 availability)',
    'Latest Unresolved Issue': 'MAIN .51 unavailable for centralized serial validation; NODE_02/NODE_03 ESP32 hardware not yet assembled.',
    'Testing State': 'Firmware: 4-env build PASS. RPi: py_compile PASS. Tests: test_api_devices 7-logic PASS, test_lora_pipeline PASS. Pending: RPi live tests.',
    'Next Incomplete Task': 'Wire ESP32 to NODE03 GPIO UART; verify NODE01/NODE02 UART readiness.',
}

for row in ws24.iter_rows(min_row=1, max_row=ws24.max_row, max_col=2):
    field = row[0].value
    if field and str(field).strip() in updates_24:
        row[1].value = updates_24[str(field).strip()]
        print(f"Active Context: updated {field}")

# Sheet 36: Known Issues
ws36 = book['Known Issues']
for row in ws36.iter_rows(min_row=1, max_row=ws36.max_row):
    first_cell = str(row[0].value) if row[0].value else ''
    if 'USB microphone' in first_cell or ('chainsaw detection' in first_cell and 'untested' in first_cell):
        for c in row:
            if c.value and isinstance(c.value, str) and 'Install USB microphone' in c.value:
                c.value = 'Device enumeration v2 fixed (non-zero arecord -l authoritative). Detector UX, sample rate, visualizer implemented. Pending: install USB mic on a node and run live field validation.'
                print(f"Known Issues: updated I020")

# Sheet 40: Change Log
ws40 = book['Change Log']
last_row_40 = ws40.max_row
while last_row_40 > 1:
    has_data = any(ws40.cell(row=last_row_40, column=c).value is not None for c in range(1, 8))
    if has_data:
        break
    last_row_40 -= 1
next_row_40 = last_row_40 + 1

new_changes = [
    ['C029', '2026-06-09', 'Dashboard', 'Separate NODE dashboard for .52/.53/.54; MAIN preserved; simulation forced to live', 'node_dashboard.html; node_app.js; app.py; config.json', 'py_compile PASS', 'Various', 'NODE GUI cleanup'],
    ['C030', '2026-06-10', 'Detection', 'Chainsaw detector UX: device selection, error visibility, Starting feedback, audio visualizer', 'detector.py; app.py; node_app.js', 'py_compile PASS', 'Various', 'Field validation'],
    ['C031', '2026-06-10', 'Detection', 'Sample rate fix: detector uses device default_samplerate (44100 Hz)', 'detector.py', 'Detector starts without PaError', 'a1907e3', 'Test with other USB mic models'],
    ['C032', '2026-06-10', 'Detection', 'Live device enumeration: /api/devices uses arecord -l for ALSA query', 'app.py', 'Empty arecord -> empty device list', '2ab870f', 'Field validation'],
    ['C033', '2026-06-12', 'Detection', 'Live device enumeration v2: non-zero arecord exit authoritative; regression tests', 'app.py; tests/test_api_devices.py', '7 logic tests PASS', '423fb63', 'RPi live validation'],
    ['C034', '2026-06-12', 'Firmware', 'Multi-node firmware: 4 PlatformIO envs (main/node_01/node_02/node_03)', 'main.cpp; platformio.ini', 'All 4 envs SUCCESS', 'daf78c4', 'Flash when ESP32 arrives'],
]

for i, change in enumerate(new_changes):
    for j, val in enumerate(change):
        ws40.cell(row=next_row_40 + i, column=j + 1, value=val)

print(f"Change Log: added {len(new_changes)} entries")

# Sheet 41: Serial Update
ws41 = book['2026-06-08 Serial Update']
last_row_41 = ws41.max_row
while last_row_41 > 1:
    has_data = any(ws41.cell(row=last_row_41, column=c).value is not None for c in range(1, 5))
    if has_data:
        break
    last_row_41 -= 1
next_row_41 = last_row_41 + 1

serial_updates = [
    ['17', 'Software', 'Live device enumeration v2', 'FIXED', 'arecord -l non-zero exit now authoritative; no stale PortAudio. tests/test_api_devices.py added.'],
    ['18', 'Firmware', 'Multi-node firmware builds', 'DONE', '4 PlatformIO envs produce distinct firmware.bin. Flash: 21.5%, RAM: 6.6%.'],
    ['19', 'Software', 'NODE03 UART readiness', 'VERIFIED', '/dev/serial0 exists; serial reader connected; local_serial.connected=true.'],
    ['20', 'Hardware', 'NODE_02/NODE_03 ESP32', 'PENDING', 'Firmware ready. ESP32 + LoRa + sensor modules not yet assembled.'],
    ['21', 'Hardware', 'MAIN .51 availability', 'BLOCKED', '.51 unavailable. Centralized serial validation and multi-node integration blocked.'],
]

for i, update in enumerate(serial_updates):
    for j, val in enumerate(update):
        ws41.cell(row=next_row_41 + i, column=j + 1, value=val)

print(f"Serial Update: added {len(serial_updates)} entries")

book.save(wb_path)
print(f"\nWorkbook saved: {wb_path}")
print("Sync complete.")
